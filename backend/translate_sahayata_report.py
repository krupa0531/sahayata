from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'Sahayata_Technical_and_Bank_Loan_Workflow_Report.docx'
TARGET = ROOT / 'outputs' / 'Sahayata_Technical_and_Bank_Loan_Workflow_Report_English.docx'

PARAGRAPHS = {
0: 'SAHAYATA',
1: 'Technical Stack, Problem-Solution, Machine Learning and Bank Loan Workflow',
2: 'Product focus: An alternative-credit platform for gig workers, delivery partners, street vendors and daily-wage workers.',
3: 'Prepared from: Current backend source code and README',
4: 'Document purpose: Project explanation, demo/presentation and discussion with banks or NBFCs',
5: '1. Project overview',
6: 'Sahayata is a backend-driven financial-support platform designed to help workers access formal credit when they do not have regular salary slips, income-tax returns, collateral or a strong CIBIL history. The platform derives cash-flow signals from UPI transaction history, estimates risk, flags suspicious transactions and proposes repayment plans aligned with daily income cycles.',
7: 'Important: the repository credit models are trained on demo or synthetic data. The current output must not be treated as a final lending decision. In production, lender policy, consented data, compliance checks, validation and human review are essential.',
8: '2. Problem being solved',
10: '3. Technology stack',
12: '4. How machine learning is used',
13: '4.1 Feature engineering: insights derived from UPI data',
14: 'Daily throughput: the average UPI credit or inflow for each active day.',
15: 'Transaction volume: the average number of unique customers per day; a proxy for customer footfall.',
16: 'Consistency score: transacting days divided by total observed days, multiplied by 100.',
17: 'Customer retention: repeat customers divided by total unique customers, multiplied by 100.',
18: 'Average monthly revenue: observed UPI inflows scaled to a 30-day monthly estimate.',
19: 'Weekday sales profile: average collection by day of week, used to adjust flexible repayment recommendations.',
20: '4.2 Supervised credit-risk model',
21: 'The implementation in ml/credit_scorer.py uses RandomForestClassifier. The input vector includes daily throughput, transaction volume, consistency score, customer retention and average monthly revenue. The model classifies the risk level as low, medium, high or reject.',
23: 'Formula: Maximum indicative loan eligibility = average monthly revenue x risk multiplier.',
24: '4.3 Fraud and anomaly detection',
25: 'The implementation in ml/anomaly_detector.py uses IsolationForest. It checks high-value inflows as statistical anomalies and against a 5x daily-average spike rule. A same-day outflow of approximately 80% after an INR 50,000 or higher credit may trigger a velocity-risk flag. The warning helps a bank or NBFC request source verification or manual review.',
26: '4.4 Flexible repayment logic',
27: 'The current implementation adjusts the daily deduction percentage between 5% and 18% using the weekday sales pattern. A deduction may be higher on strong weekend sales days and lower on slower days. The system returns daily-fixed, daily-flexible and monthly plans; the recommended mode is daily_flexible.',
28: '5. End-to-end loan journey',
29: 'The worker registers or signs in with a mobile number and verifies the OTP through the official flow.',
30: 'The worker completes e-KYC, bank or UPI linkage and transaction-data consent. An OTP must never be shared with an agent or assistant.',
31: 'The backend fetches UPI transactions and performs feature engineering.',
32: 'The Random Forest model produces a risk score and maximum indicative loan eligibility; the anomaly detector produces a warning list.',
33: 'The worker chooses a loan amount and repayment mode, then submits an application.',
34: 'Application timeline: Submitted -> e-KYC verified -> NBFC review -> Bank sanction -> Cash disbursed.',
35: 'After final verification, the bank or NBFC approves, rejects or modifies the application. In approved cases, funds are disbursed to the verified bank account.',
36: 'Repayments may be collected through consented eNACH, UPI Autopay or monthly auto-debit routes. A missed-payment retry and escalation policy applies.',
37: '6. What the bank does and how the loan is issued',
38: 'Sahayata is not a bank. In the current design, it functions as an origination, alternative-score, customer-journey and monitoring layer. Actual lending, sanction, pricing, disbursement and regulated collections must remain under the control of a bank or RBI-regulated NBFC.',
40: '7. Recommended bank decision rule for production',
41: 'KYC and account validation: verify identity, mobile number, bank account and required documents.',
42: 'Explicit consent: collect UPI or bank transaction data with a stated purpose, duration and revocation rights.',
43: 'Data quality: check for meaningful minimum transaction history, duplicate or fake patterns and missing data.',
44: 'Risk score and policy rules: use the ML score as one input and combine it with bureau data, affordability, exposure limits and fraud or AML rules.',
45: 'Anomaly hold: where a critical warning exists, stop automated disbursal and request source-of-funds or manual verification.',
46: 'Affordability: keep repayment within a safe share of verified disposable cash flow; show the borrower the amount, APR or interest, fees, tenure and total payable clearly.',
47: 'Human review and adverse action: provide manual review for borderline or rejected cases and give an understandable reason or improvement path.',
48: 'Disburse and monitor: after the signed agreement and mandate, transfer funds to the account, track repayment and maintain a customer grievance channel.',
49: '8. Security and compliance considerations',
50: 'JWT-based authenticated endpoints and password hashing are present. Production should add secret management, key rotation, HTTPS and audit logs.',
51: 'The user must enter an OTP on the official screen. The project guidance clearly says that OTPs must never be shared.',
52: 'UPI or bank-data access must not happen without explicit informed consent. Retain only data necessary for the stated purpose.',
53: 'ML output must be explainable, bias-tested and monitored. A synthetic demo training dataset must not be deployed as a real lending model.',
54: 'Before launch, legal and compliance teams must validate RBI and partner-bank or NBFC policy, KYC and AML requirements, data-protection obligations, fair-lending practices, grievance redressal and pricing disclosures.',
55: '9. Current implementation versus production readiness',
57: '10. Conclusion',
58: 'Sahayata has a strong core idea: instead of evaluating daily-income workers only through traditional salary documents, it can help lenders understand consented UPI cash flow, consistency and customer behaviour. Machine learning should assist decisions by identifying risk signals and repayment fit. The final loan decision must always be made by a regulated bank or NBFC with compliant underwriting and transparent customer disclosures.'
}

TABLES = [
[
 ['Worker problem', 'Sahayata approach'],
 ['Income is daily and irregular; a monthly EMI may not match the cash cycle.', 'Uses UPI inflows to understand daily cash flow and offers daily-fixed, daily-flexible and monthly repayment options.'],
 ['No salary slip, employment contract or income-tax return.', 'Builds alternative underwriting signals from consented UPI transaction history.'],
 ['New-to-credit workers may have limited collateral and CIBIL history.', 'Uses cash-flow consistency, transaction volume and repeat customers as risk inputs.'],
 ['Manual processing costs are high for small-ticket loans.', 'Uses FastAPI APIs, automated feature extraction and application tracking to standardise operations.'],
 ['Risk of fraud or unusual cash inflows.', 'Uses Isolation Forest and rule checks to flag suspicious spikes and possible rapid outflow.'],
],
[
 ['Layer', 'Technology', 'Role'],
 ['Backend API', 'Python + FastAPI + Uvicorn', 'REST APIs, request validation, loan workflow and API documentation.'],
 ['Data validation', 'Pydantic', 'Validates request and response schemas.'],
 ['Database', 'SQLite (current) + PostgreSQL schema', 'Stores users, applications, status events, UPI transactions and demo data.'],
 ['Machine learning', 'scikit-learn + NumPy', 'Random Forest credit-risk classification and Isolation Forest anomaly detection.'],
 ['Authentication', 'JWT via python-jose', 'Bearer-token login and session protection.'],
 ['Password security', 'PBKDF2-HMAC-SHA256 + random salt', 'Does not store passwords in plain text.'],
 ['OTP', 'SMS OTP service integration', 'Mobile verification; intended not to store or log the code.'],
 ['Voice guidance', 'ElevenLabs (optional) + Edge TTS fallback', 'Hindi, Gujarati and English audio guidance for loan journeys.'],
 ['Frontend integration', 'React app proxy (README reference)', 'Routes frontend /api calls to the backend.'],
],
[
 ['Risk level', 'Loan multiplier', 'Indicative outcome'],
 ['Low', '2.0x monthly revenue', 'Consistent cash flow; higher indicative eligibility.'],
 ['Medium', '1.0x monthly revenue', 'Begin with a smaller working-capital loan.'],
 ['High', '0.5x monthly revenue', 'Reduced limit and stricter monitoring.'],
 ['Reject', '0x', 'Insufficient UPI history; recommend building more active transaction days.'],
],
[
 ['Stage', 'Sahayata / FinTech role', 'Bank / NBFC role'],
 ['Onboarding', 'App experience, mobile verification, application capture and consent screen.', 'KYC policy, permitted partner controls and customer-due-diligence oversight.'],
 ['Data and score', 'Consented UPI features, indicative score, anomaly warnings and repayment recommendation.', 'Score validation, bureau, AML and fraud checks, underwriting rules and final eligibility.'],
 ['Credit decision', 'Provides an application package and explanation.', 'Final approval, rejection or modification; determines loan amount, interest or APR, tenure, fees and conditions.'],
 ['Sanction', 'Shows application status and timeline to the worker.', 'Issues sanction letter, Key Fact Statement, disclosures and agreement.'],
 ['Disbursement', 'Updates the approved-workflow status.', 'Transfers funds to the verified borrower bank account.'],
 ['Repayment', 'UPI or Autopay reminders, cash-flow-aware plan and support interface.', 'Mandate setup, collection, ledger, reconciliation, hardship handling and compliant recovery.'],
 ['Monitoring', 'Fresh UPI patterns and anomaly alerts, subject to consent and policy.', 'Portfolio monitoring, delinquency handling, reporting, audit and regulatory compliance.'],
],
[
 ['Current code', 'Next step for production'],
 ['Random Forest demo model on synthetic labelled profiles.', 'Use consented historical repayment labels, a validation set, bias and fairness testing, model monitoring and governance.'],
 ['SQLite demo database and seeded demo transactions.', 'Use an encrypted managed database, access controls, backup, retention/deletion policy and audit trail.'],
 ['Rule-based anomaly checks plus Isolation Forest.', 'Use a calibrated fraud strategy, case-management queue, source verification and alert-performance measurement.'],
 ['Simple repayment calculation.', 'Use a regulated loan schedule, interest or APR computation, Key Fact Statement, mandate lifecycle, reconciliation and hardship policy.'],
 ['Application status tracking.', 'Integrate with bank or NBFC LOS/LMS, bureau, AML and KYC connectors and human underwriting workflow.'],
]
]

def replace_paragraph(paragraph, text):
    # Retain the original paragraph style and alignment while replacing mixed-language copy.
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)

def replace_cell(cell, text):
    p = cell.paragraphs[0]
    replace_paragraph(p, text)
    for other in cell.paragraphs[1:]:
        other._element.getparent().remove(other._element)

def main():
    doc = Document(SOURCE)
    for idx, text in PARAGRAPHS.items():
        replace_paragraph(doc.paragraphs[idx], text)
    for table, rows in zip(doc.tables, TABLES):
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement('w:tblHeader')
        header.set(qn('w:val'), 'true')
        tr_pr.append(header)
        for row, source_row in zip(table.rows, rows):
            for cell, value in zip(row.cells, source_row):
                replace_cell(cell, value)
    doc.core_properties.title = 'Sahayata Technical and Bank Loan Workflow Report'
    doc.core_properties.subject = 'English translation'
    doc.core_properties.author = 'Sahayata Team'
    doc.save(TARGET)
    print(TARGET)

if __name__ == '__main__':
    main()
