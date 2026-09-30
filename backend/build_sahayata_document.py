from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
OUT.mkdir(exist_ok=True)
ASSETS = OUT / "sahayata_doc_assets"
ASSETS.mkdir(exist_ok=True)
DOCX = OUT / "Sahayata_Judge_Pitch_and_Technical_Documentation.docx"

NAVY = "0B2545"; BLUE = "2E74B5"; TEAL = "0F766E"; GOLD = "A16207"; LIGHT = "E8EEF5"; PALE = "F4F6F9"; GRAY = "5B6770"; RED = "9B1C1C"; GREEN = "166534"

def font(size, bold=False):
    try: return ImageFont.truetype("arialbd.ttf" if bold else "arial.ttf", size)
    except: return ImageFont.load_default()

def center(draw, box, text, f, fill, spacing=4):
    l,t,r,b = draw.multiline_textbbox((0,0), text, font=f, spacing=spacing, align="center")
    draw.multiline_text(((box[0]+box[2]-(r-l))/2, (box[1]+box[3]-(b-t))/2), text, font=f, fill=fill, spacing=spacing, align="center")

def rounded(draw, xy, fill, outline=None, radius=20, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)

def arrow(draw, a, b, fill="#5B6770", width=4):
    draw.line([a,b], fill=fill, width=width)
    x,y=b; draw.polygon([(x,y),(x-12,y-7),(x-12,y+7)], fill=fill)

def architecture_image():
    p=ASSETS/"architecture.png"; im=Image.new("RGB",(1700,1240),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1700,76),"SAHAYATA | MERMAID-STYLE PRODUCTION ARCHITECTURE",font(33,True),"#0B2545")
    # External actors
    nodes=[(55,150,310,290,"Worker\nMobile / Voice","#E7F1FF"),(55,450,310,590,"Bank Officer\nCommand Center","#F2EEFF"),(1390,280,1645,420,"Regulated Lender\n/ AA / KYC Partner","#FFF7E6")]
    for x1,y1,x2,y2,t,c in nodes: rounded(d,(x1,y1,x2,y2),c,"#7E9AB6",22,2); center(d,(x1+15,y1+20,x2-15,y2-20),t,font(23,True),"#0B2545")
    # Platform bands
    bands=[(380,110,1320,300,"EXPERIENCE & ORCHESTRATION","Worker app | multilingual SAI | consent UX | API gateway","#EAF8F4"),(380,360,1320,610,"INTELLIGENCE & DECISIONING","Policy rules | credit features | fraud signals | source-cited scheme guidance","#E7F1FF"),(380,670,1320,915,"DATA & OPERATIONS","Encrypted documents | application workflow | underwriting queue | lender adapter","#F2EEFF"),(380,980,1320,1145,"TRUST FOUNDATION","RBAC | audit ledger | monitoring | retention policy | encryption","#F4F6F9")]
    for x1,y1,x2,y2,h,sub,c in bands:
      rounded(d,(x1,y1,x2,y2),c,"#8CA5BF",26,2); d.text((420,y1+25),h,font=font(26,True),fill="#0B2545")
    # Components inside bands
    comps=[(430,175,635,270,"Voice + Text\nWorker Portal"),(735,175,940,270,"SAI\nOrchestrator"),(1040,175,1265,270,"Consent\nManager"),(430,445,635,550,"Credit\nEngine"),(735,445,940,550,"Rules + Scheme\nKnowledge"),(1040,445,1265,550,"Fraud\nEngine"),(430,755,635,860,"Document\nVault"),(735,755,940,860,"Application\nService"),(1040,755,1265,860,"Lender\nAdapter")]
    for x1,y1,x2,y2,t in comps: rounded(d,(x1,y1,x2,y2),"white","#2E74B5",14,2); center(d,(x1+8,y1+8,x2-8,y2-8),t,font(18,True),"#0B2545")
    for a,b in [((310,220),(375,220)),((635,222),(730,222)),((940,222),(1035,222)),((1265,807),(1385,350)),((310,520),(375,520)),((635,497),(730,497)),((940,497),(1035,497)),((532,550),(532,750)),((837,550),(837,750)),((1142,550),(1142,750))]: arrow(d,a,b,"#2E74B5",4)
    center(d,(390,1160,1310,1215),"Final approval, pricing and disbursal remain with the regulated lender. Sahayata recommends, explains and routes.",font(20,True),"#9B1C1C")
    im.save(p); return p

def modules_image():
    p=ASSETS/"modules.png"; im=Image.new("RGB",(1700,930),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1700,76),"ONE PLATFORM. FOUR CONNECTED MODULES.",font(33,True),"#0B2545")
    cards=[(75,140,805,420,"01 | WORKER EXPERIENCE","Voice-first journeys\nHindi | Gujarati | English\nKYC uploads | scheme discovery","#E7F1FF"),(895,140,1625,420,"02 | SAI ASSISTANT","Guided intake\nClear consent prompts\nSource-cited financial guidance","#EAF8F4"),(75,505,805,785,"03 | TRUSTED DECISIONING","Cashflow features | affordability\nFraud investigation signals\nReason codes + human review","#FFF7E6"),(895,505,1625,785,"04 | LENDER OPERATIONS","Structured applications\nUnderwriting queue\nPartner integrations + repayment support","#F2EEFF")]
    for x1,y1,x2,y2,h,sub,c in cards:
      rounded(d,(x1,y1,x2,y2),c,"#8CA5BF",28,2); d.text((x1+40,y1+40),h,font=font(27,True),fill="#0B2545"); d.multiline_text((x1+40,y1+115),sub,font=font(23),fill="#35424E",spacing=13)
    d.line((850,420,850,505),fill="#2E74B5",width=7); d.line((805,280,895,280),fill="#2E74B5",width=7); d.line((805,645,895,645),fill="#2E74B5",width=7)
    center(d,(100,830,1600,895),"The worker sees one simple journey; the platform coordinates data, policy, security and lender operations behind the scenes.",font(22,True),"#0F766E")
    im.save(p); return p

def tech_flow_image():
    p=ASSETS/"ai_tech_flow.png"; im=Image.new("RGB",(1700,1160),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1700,76),"AI DECISION FLOW | FROM CONSENT TO LENDER-READY CASE",font(31,True),"#0B2545")
    steps=[("1. Permission","Consent scope\n+ purpose"),("2. Evidence","UPI / statement\n+ document input"),("3. Preparation","OCR + parser\n+ quality checks"),("4. Features","Income regularity\nexpenses + buffer"),("5. Decision support","Credit + fraud\n+ policy rules"),("6. Governance","Reason codes\n+ human review"),("7. Partner action","Lender decision\n+ disbursal")]
    xs=[40,275,510,745,980,1215,1450]
    for i,(h,sub) in enumerate(steps):
      rounded(d,(xs[i],195,xs[i]+205,470),"#F4F6F9","#2E74B5",18,2); center(d,(xs[i]+12,225,xs[i]+193,290),h,font(19,True),"#0B2545"); center(d,(xs[i]+15,320,xs[i]+190,425),sub,font(20),"#485662")
      if i<6: arrow(d,(xs[i]+205,333),(xs[i+1]-8,333),"#2E74B5",3)
    # Gates
    gates=[(125,620,510,780,"QUALITY GATE","Missing or inconsistent evidence?\nAsk worker for clarity or route to review."),(655,620,1045,780,"SAFETY GATE","High fraud risk or low confidence?\nHold automation; notify investigator."),(1190,620,1580,780,"LENDER GATE","Licensed partner owns final\napproval, price and disbursal.")]
    for x1,y1,x2,y2,h,sub in gates:
      rounded(d,(x1,y1,x2,y2),"#FFF7E6","#A16207",18,2); center(d,(x1+15,y1+25,x2-15,y1+70),h,font(21,True),"#7A5A00"); center(d,(x1+15,y1+85,x2-15,y2-15),sub,font(19),"#35424E")
    for x in [300,850,1400]: arrow(d,(x,470),(x,610),"#A16207",4)
    rounded(d,(80,875,1620,1050),"#EAF8F4","#0F766E",20,2)
    center(d,(110,900,1590,1025),"AUDIT EVENT CREATED AT EVERY STEP\nconsent | data access | model/policy version | alert | officer action | customer communication",font(24,True),"#0B2545")
    im.save(p); return p

def trust_image():
    p=ASSETS/"trust.png"; im=Image.new("RGB",(1700,700),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1700,76),"TRUST IS A PRODUCT FEATURE, NOT A FOOTNOTE",font(31,True),"#0B2545")
    values=[("CONSENT","Purpose, scope, expiry\nand revocation"),("SECURITY","Encryption, RBAC,\nsecret management"),("EXPLAINABILITY","Reason codes, source links\nand clear next actions"),("ACCOUNTABILITY","Audit trail, human review\nand correction path")]
    xs=[65,480,895,1310]
    for x,(h,sub) in zip(xs,values):
      d.ellipse((x+120,155,x+240,275),fill="#2E74B5"); center(d,(x+120,155,x+240,275),"✓",font(55,True),"white"); rounded(d,(x,300,x+360,560),"#F4F6F9","#8CA5BF",22,2); center(d,(x+25,335,x+335,385),h,font(23,True),"#0B2545"); center(d,(x+20,415,x+340,515),sub,font(21),"#485662")
    im.save(p); return p

def workflow_image():
    p=ASSETS/"workflow.png"; im=Image.new("RGB",(1600,940),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1600,75),"WORKER-TO-LOAN JOURNEY",font(31,True),"#0B2545")
    steps=[("1","Talk to SAI","Voice or text\nin your language"),("2","Give consent","Choose permitted\nfinancial data"),("3","Check eligibility","Income pattern +\naffordability"),("4","Risk review","Fraud signals +\nhuman review"),("5","Apply to lender","Structured case\nto partner"),("6","Flexible repayment","Mandate only after\nclear consent")]
    xs=[75,595,1115,75,595,1115]
    for i,(n,h,sub) in enumerate(steps):
       row=0 if i<3 else 1; y=145 if row==0 else 470
       rounded(d,(xs[i],y,xs[i]+410,y+240),"#F4F6F9","#8CA5BF",20,2)
       d.ellipse((xs[i]+172,y+24,xs[i]+238,y+90),fill="#2E74B5"); center(d,(xs[i]+172,y+24,xs[i]+238,y+90),n,font(28,True),"white")
       center(d,(xs[i]+20,y+112,xs[i]+390,y+157),h,font(27,True),"#0B2545")
       center(d,(xs[i]+18,y+172,xs[i]+392,y+222),sub,font(21),"#485662")
       if i in [0,1,3,4]: arrow(d,(xs[i]+410,y+120),(xs[i+1]-12,y+120),"#2E74B5",4)
    arrow(d,(1320,385),(1320,440),"#2E74B5",4)
    center(d,(35,790,1565,860),"Safe design: no OTP collection by the assistant | no automatic approval | no sale of raw personal financial data.",font(23,True),"#9B1C1C")
    im.save(p); return p

def data_image():
    p=ASSETS/"inclusion_data.png"; im=Image.new("RGB",(1600,800),"white"); d=ImageDraw.Draw(im)
    center(d,(0,20,1600,72),"THE FINANCIAL-INCLUSION CONTEXT",font(31,True),"#0B2545")
    cards=[("~90%","informally\nemployed","ILO India Employment\nReport 2024"),("30.68 crore","e-Shram workers\nregistered","Ministry of Labour\nMar 2025"),("35 million","first credit product\nopened in 2021","TransUnion CIBIL"),("27.18%","met all financial-\nliteracy thresholds","NCFE national survey\n2019")]
    x=65
    for n,label,source in cards:
      rounded(d,(x,150,x+330,545),"#F4F6F9","#B5C5D6",20,2)
      center(d,(x+20,195,x+310,290),n,font(37,True),"#0F766E")
      center(d,(x+27,315,x+303,415),label,font(24,True),"#0B2545")
      center(d,(x+25,440,x+305,510),source,font(19),"#5B6770")
      x+=385
    d.line((65,610,1535,610), fill="#B5C5D6",width=2)
    d.text((65,645),"No definitive nationwide count exists for people with no CIBIL score; ‘new-to-credit’ is a transparent proxy.",font=font(22),fill="#35424E")
    d.text((65,700),"These figures show the scale of the gap; they do not prove an individual worker’s creditworthiness.",font=font(22,True),fill="#9B1C1C")
    im.save(p); return p

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def set_cell_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(dxa)); tcW.set(qn('w:type'),'dxa')
def set_table_geometry(table, widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr; tblW=tblPr.first_child_found_in('w:tblW'); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    for row in table.rows:
      for cell,w in zip(row.cells,widths): set_cell_width(cell,w); set_cell_margin(cell)

def set_run(run, size=11, color=NAVY, bold=False, italic=False):
    run.font.name='Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color); run.bold=bold; run.italic=italic
def add_para(doc, text='', style=None, size=11, color=NAVY, bold=False, italic=False, align=None, after=8):
    # Convenience: allow add_para(doc, text, 14, COLOR, ...) for short display text.
    if isinstance(style, (int, float)):
        if isinstance(size, str):
            color = size
        size, style = style, None
    p=doc.add_paragraph(style=style); p.paragraph_format.space_after=Pt(after)
    if align: p.alignment=align
    r=p.add_run(text); set_run(r,size,color,bold,italic); return p
def add_bullets(doc, items):
    for item in items:
       p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.2
       set_run(p.add_run(item),11,NAVY)
def add_heading(doc, text, level=1): doc.add_heading(text,level)
def add_table(doc, headers, rows, widths):
    table=doc.add_table(rows=1,cols=len(headers)); table.style='Table Grid'; set_table_geometry(table,widths)
    for cell,text in zip(table.rows[0].cells,headers):
      set_cell_shading(cell,LIGHT); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
      p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); set_run(p.add_run(text),10,NAVY,True)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
    for row in rows:
      cells=table.add_row().cells
      for cell,text in zip(cells,row):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); set_run(p.add_run(text),9.5,NAVY)
    return table
def add_caption(doc,text):
    p=add_para(doc,text,size=9,color=GRAY,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=12); return p

def setup(doc):
  sec=doc.sections[0]; sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
  styles=doc.styles
  normal=styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(NAVY); normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.25
  for name,size,color,before,after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,NAVY,8,4)]:
    s=styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
  header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_run(header.add_run('SAHAYATA | Judge Pitch & Technical Documentation'),9,GRAY)
  footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(footer.add_run('Sahayata | Confidential prototype concept | August 2026'),9,GRAY)

def page_break(doc): doc.add_page_break()

def build():
  arch=architecture_image(); flow=workflow_image(); context=data_image(); modules=modules_image(); tech_flow=tech_flow_image(); trust=trust_image()
  doc=Document(); setup(doc)
  # Cover
  p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(92); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  set_run(p.add_run('SAHAYATA'),34,NAVY,True)
  p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(p.add_run('Financial inclusion infrastructure for India’s informal and gig workers'),19,TEAL,True)
  add_para(doc,'Judge Pitch & Technical Feature Documentation',14,GRAY,align=WD_ALIGN_PARAGRAPH.CENTER,after=34)
  t=doc.add_table(rows=1,cols=3); t.style='Table Grid'; set_table_geometry(t,[3120,3120,3120])
  trPr = t.rows[0]._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
  for cell,(n,l) in zip(t.rows[0].cells,[('Voice-first','Guidance in Hindi, Gujarati and English'),('Alternative data','Consent-led cashflow signals, not a replacement for regulated bureau data'),('Trust by design','Human review, audit trails and lender-controlled decisions')]):
    set_cell_shading(cell,PALE); p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(p.add_run(n+'\n'),14,NAVY,True); set_run(p.add_run(l),9.5,GRAY)
  add_para(doc,'Prepared for evaluation and demonstration',11,GRAY,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
  add_para(doc,'Version 1.0 | Research cut-off: August 2026',10,GRAY,align=WD_ALIGN_PARAGRAPH.CENTER)
  page_break(doc)
  # Executive
  add_heading(doc,'1. Executive summary',1)
  add_para(doc,'Sahayata is a voice-first financial-inclusion platform concept for informal and gig workers. It helps workers understand government schemes, submit a consent-based financial profile, receive explainable eligibility guidance, and apply through regulated lender partners. The system is deliberately designed as a decision-support and origination layer: it does not promise a loan, make autonomous final credit decisions, or collect OTPs.')
  add_heading(doc,'The problem in one minute',2)
  add_para(doc,'Many workers have recurring earnings but lack salary slips, formal employment records, or a deep bureau history. Existing processes can be hard to navigate in a worker’s preferred language, and scheme information is fragmented. Sahayata combines guided onboarding, consent-aware cashflow analysis, scheme discovery and risk controls to make the path to formal finance easier to understand and safer to operate.')
  add_heading(doc,'What makes it different',2)
  add_bullets(doc,['Voice and text journeys in Hindi, Gujarati and English for lower-friction onboarding.','Alternative-data features based on consented transaction patterns and affordability, with clear reason codes.','A policy knowledge layer that cites reviewed official sources instead of inventing scheme advice.','Fraud signals and a human-underwriting escalation path before a lender acts.','Flexible repayment-plan illustration; the regulated lender determines final product, pricing and mandate terms.'])
  add_heading(doc,'Judge takeaway',2)
  add_para(doc,'Sahayata turns a difficult, paper-heavy journey into a simple sequence: explain → consent → assess → review → route. Its value is not “AI approves loans”; its value is helping workers and lenders make safer, more informed decisions.',size=11,color=TEAL,bold=True)
  # Evidence
  page_break(doc); add_heading(doc,'2. Evidence: why this problem matters',1); doc.add_picture(str(context),width=Inches(6.85)); add_caption(doc,'Figure 1. India’s inclusion context. Sources and definitions are listed in Section 12.')
  add_heading(doc,'Interpreting the data responsibly',2)
  add_para(doc,'The evidence supports a large informal-workforce and financial-capability challenge. It does not support a claim that every informal worker is unbanked or has no CIBIL score. India has made important progress in access: RBI’s FI-Index reached 67.0 for March 2025. Sahayata is designed for the remaining usability, documentation and “thin-file” gaps—not as a substitute for the banking system.')
  add_heading(doc,'Research-backed problem statements',2)
  add_table(doc,['Observed gap','Evidence','Design response'],[
   ('Informal and irregular income','ILO reports nearly 90% of the workforce is informally employed.','Use consented cashflow consistency and affordability signals alongside lender policy.'),
   ('Thin or new credit files','TransUnion CIBIL reported 35 million consumers opened their first credit product in 2021; 67% of that segment was rural/semi-urban.','Do not reject solely for lack of traditional history; route to appropriate lender policy and document reasons.'),
   ('Financial capability gap','NCFE’s 2019 survey found 27.18% met the minimum threshold across all three literacy components.','Explain products, pricing, repayment and grievance routes in simple language and voice.'),
   ('Scheme discoverability','RBI and government programmes maintain extensive financial-literacy and outreach infrastructure; awareness remains an explicit policy priority.','Provide source-cited scheme discovery, never an unsupported guarantee of eligibility.')], [2100,3100,4160])
  # Solution
  page_break(doc); add_heading(doc,'3. Sahayata solution and product modules',1)
  doc.add_picture(str(modules),width=Inches(6.85)); add_caption(doc,'Figure 2. The product is deliberately organised into four connected modules so the worker experience remains simple.')
  add_heading(doc,'A. Worker portal and financial tools',2); add_bullets(doc,['Loan-eligibility calculator based on worker-entered earnings, expenses and household needs.','Document upload and status tracking for KYC or statement evidence.','Scheme directory for PM SVANidhi, PM-SYM, PMSBY, PMJJBY and PMJDY, backed by official-source references.','Plain-language safety prompts: never share OTP, PIN, CVV or remote-device access.'])
  add_heading(doc,'B. SAI: Sahayata AI assistant',2); add_bullets(doc,['Collects information step-by-step by voice or text; supports multilingual prompts.','Explains why a data point is requested and captures consent before data access.','Provides an eligibility report with assumptions, missing evidence and next actions.','Answers policy questions only from a reviewed knowledge base and returns source citations.'])
  add_heading(doc,'C. Credit and affordability intelligence',2); add_bullets(doc,['Feature examples: transaction frequency, income regularity, inflow volatility, retained cash buffer and anomaly indicators.','Outputs a recommendation band and reason codes, not an autonomous approval.','Final underwriting policy, loan offer and disbursal remain with the licensed lender.'])
  add_heading(doc,'D. Fraud and underwriting operations',2); add_bullets(doc,['Document intake can flag metadata anomalies, OCR mismatch, duplicate submissions and suspicious transaction patterns.','High-risk or low-confidence cases are placed in a manual-review queue.','Every action is timestamped for auditability: consent, document access, model output, officer decision and customer notification.'])
  # journey
  page_break(doc); add_heading(doc,'4. Worker journey: simple enough to demonstrate',1); doc.add_picture(str(flow),width=Inches(6.85)); add_caption(doc,'Figure 3. The end-to-end journey. A worker can stop after any step; consent can be revoked subject to legal retention obligations.')
  add_heading(doc,'Demo scenario for judges',2)
  add_para(doc,'Example: a 28-year-old street vendor from Ahmedabad uses Hindi or Gujarati. SAI asks about average earnings and essential expenses, explains the requested bank/UPI data, and shows an affordability estimate. The worker receives a source-cited scheme shortlist and, if they choose, submits a lender application. A flagged document or unusual cashflow pattern moves the case to an officer—not an automatic rejection.')
  add_heading(doc,'What the system must never do',2)
  add_bullets(doc,['Ask for an OTP, UPI PIN, CVV or banking password.','Claim that a government-scheme benefit or loan is approved before the authorised entity confirms it.','Treat a fraud score as proof of wrongdoing.','Use a score as the only reason for a final lending decision.','Retain raw documents longer than the documented purpose and retention policy allow.'])
  # architecture
  page_break(doc); add_heading(doc,'5. Production architecture',1); doc.add_picture(str(arch),width=Inches(6.85)); add_caption(doc,'Figure 4. Mermaid-style target architecture. The prototype can start modular and split services as integrations and volume grow.')
  add_heading(doc,'Architecture principles',2)
  add_table(doc,['Principle','Implementation implication'],[
   ('Consent first','Capture purpose, scope, timestamp, expiry and revocation; only ingest data allowed by the active consent.'),
   ('Separation of duties','Keep worker experience, decisioning, lender integration and audit functions logically separated.'),
   ('Explainability','Store the features, policy version, result and reason codes shown for each recommendation.'),
   ('Fail safely','If a provider, model or verification step fails, pause or route to review; never silently approve.'),
   ('Least privilege','Encrypt sensitive fields; use RBAC, staff MFA, environment-specific secrets and audited access.')], [2400,6960])
  add_heading(doc,'Suggested deployment evolution',2)
  add_bullets(doc,['Prototype: FastAPI, PostgreSQL, object storage, background jobs, isolated ML modules and reviewed-source RAG.','Pilot: API gateway/WAF, Redis queue, encrypted object storage, observability, consent ledger and lender sandbox adapters.','Production: service boundaries for identity, consent/data, decisioning, lending operations and trust; disaster recovery, penetration testing and third-party risk review.'])
  # technical
  page_break(doc); add_heading(doc,'6. Technical feature blueprint',1); doc.add_picture(str(tech_flow),width=Inches(6.85)); add_caption(doc,'Figure 5. AI-assisted decision flow. “AI” supports preparation and triage; the partner lender owns the credit decision.')
  add_table(doc,['Component','Prototype capability','Production hardening'],[
   ('Identity and access','Mobile OTP, JWT session, authenticated worker endpoints.','OTP vendor controls, MFA for staff, device/session risk, rate limits and breach monitoring.'),
   ('Document service','Upload KYC and statements; virus scan; encrypted storage.','Content-type validation, malware sandboxing, retention/deletion workflow and access logging.'),
   ('Credit engine','Cashflow aggregation, consistency features, repayment-plan calculation.','Feature registry, model monitoring, bias testing, champion/challenger controls and lender policy versioning.'),
   ('Fraud engine','Anomaly detection and document consistency checks.','Calibrated rules, investigator feedback loop, false-positive review and provider verification APIs.'),
   ('Knowledge service','Reviewed RBI/scheme/app-policy sources with citations.','Content approval workflow, effective dates, owner, expiry and retrieval-evaluation tests.'),
   ('Operations console','Application timeline, alerts and manual actions.','Maker-checker controls, reason-code enforcement, immutable audit export and SLA dashboards.')], [1850,3600,3910])
  add_heading(doc,'Model governance',2)
  add_para(doc,'The score should be framed as a decision-support signal. Before any real lending use, Sahayata needs documented training data provenance, performance by customer segment, reject-inference policy, drift monitoring, adverse-action reason mapping, escalation paths and periodic independent validation. A demo Random Forest or Isolation Forest is not sufficient evidence for a production lending model.')
  # Fraud
  page_break(doc); add_heading(doc,'7. Fraud controls: a responsible 8-signal approach',1)
  add_para(doc,'The original “8-layer” concept is useful as a checklist, but the correct production framing is “signals for investigation,” not a claim that every signal is accurate. Some checks require consent, vendor capability, legal review and high-quality inputs.')
  add_table(doc,['Signal family','Example check','Action'],[
   ('Document integrity','Metadata inconsistencies, duplicate image hash, OCR-field mismatch.','Request a clearer source or route to review.'),
   ('Identity verification','Official KYC provider response where permitted; selfie/liveness only with explicit consent.','Pause application when verification is inconclusive.'),
   ('Transaction anomaly','Unusual inflow spikes, rapid cash-out pattern, repeated account reuse.','Generate a reasoned alert; do not infer fraud from one anomaly.'),
   ('Device/session risk','Velocity, impossible travel, repeated device/account combinations.','Add step-up verification or manual review.'),
   ('Geo consistency','Location conflicts with worker-declared service area where legitimate and consented.','Treat as an investigation signal, not a rejection rule.'),
   ('Historical abuse','Confirmed internal cases and legally permitted, governed lists.','Apply strict access, review and correction procedures.'),
   ('Partner verification','Lender/AA/provider response and data-integrity status.','Block only on a verified provider or policy rule.'),
   ('Human review','Officer examines evidence and records outcome.','Close with decision, rationale and worker communication.')], [2050,3900,3410])
  add_heading(doc,'Fairness and customer protection',2)
  add_bullets(doc,['Use proportional data: request only what is necessary for the stated purpose.','Provide a clear path to correct profile or document errors.','Track false positives and disparate outcomes across relevant segments.','Use fraud flags to trigger review, not humiliation or opaque denial.','Ensure collection and repayment communication follows applicable lender policy and law.'])
  # Schemes
  page_break(doc); add_heading(doc,'8. Government scheme guidance in the app',1)
  add_para(doc,'Sahayata should present scheme information as discovery and pre-screening. Official portal rules, local implementation and authorised entities determine final eligibility and benefit. Rules must carry a source URL, effective date and reviewer approval.')
  add_table(doc,['Scheme','What Sahayata can explain','Source of truth'],[
   ('PM SVANidhi','Working-capital loan journey for eligible street vendors; successive loan tranches and timely-repayment incentive concepts.','Ministry of Housing and Urban Affairs / PM SVANidhi portal.'),
   ('PM-SYM','Unorganised-worker pension eligibility concepts, including age and income conditions.','Ministry of Labour and Employment / PM-SYM portal.'),
   ('PMJDY','Account services and overdraft facility information, subject to bank rules and eligibility.','Department of Financial Services / PMJDY portal.'),
   ('PMSBY / PMJJBY','Accident and life insurance scheme summaries, premiums and enrolment path.','Department of Financial Services and insurer/bank channels.'),
   ('MUDRA','Loan category information and application routing for qualifying micro enterprises.','MUDRA / participating lender guidance.')], [1550,4750,3060])
  add_heading(doc,'Information quality controls',2); add_bullets(doc,['No static, unversioned scheme rules in frontend code.','Show “last reviewed” date and official link beside every scheme recommendation.','Use a policy editor approval workflow for updates.','If the rule is uncertain, state “please verify with the official authority” rather than guess.'])
  # roadmap business
  page_break(doc); add_heading(doc,'9. Pilot plan, success metrics and business model',1)
  add_heading(doc,'90-day pilot roadmap',2)
  add_table(doc,['Phase','Objective','Deliverable'],[
    ('Weeks 1-3','Safety and data foundations','Consent UX, source-reviewed scheme directory, audit events, secure document flow.'),
    ('Weeks 4-6','Worker journey validation','Hindi/Gujarati usability tests, explanation screens, worker-support playbook.'),
    ('Weeks 7-9','Decisioning sandbox','Synthetic or approved partner data, feature validation, manual-review queue.'),
    ('Weeks 10-12','Partner pilot readiness','Lender sandbox adapter, monitoring, incident runbook, independent security review.')], [1500,3300,4560])
  add_heading(doc,'Metrics that judges and partners should ask for',2); add_bullets(doc,['Completion rate from consent to submitted application.','Time-to-decision and share of cases correctly routed to human review.','Document rework rate and fraud-alert precision/recall after investigator confirmation.','Worker comprehension: can users explain repayment amount, timing and grievance route?','Approval, default and repayment outcomes only after an adequately sized and governed lender pilot.','Fairness, complaints, consent revocations and support-resolution time.'])
  add_heading(doc,'Commercial model: validate, do not assume',2)
  add_para(doc,'Potential revenue paths include lender-originations or workflow SaaS fees, implementation fees for regulated partners, and carefully governed distribution commissions where permitted. The product should not rely on hidden worker fees, lead sale without consent, or fees for government-scheme discovery. Unit economics and regulatory suitability must be validated with each partner.')
  # Pitch
  page_break(doc); add_heading(doc,'10. Four-minute judge pitch',1)
  add_table(doc,['Time','What to show','What to say'],[
   ('0:00-0:40','Problem slide and inclusion evidence','“India has made real financial-inclusion progress, but informal workers can still face documentation, language and thin-credit-file barriers. Sahayata makes the next step easier and safer.”'),
   ('0:40-1:35','SAI in Hindi/Gujarati','“A worker speaks or types naturally. We explain each question and take consent before any financial data is used.”'),
   ('1:35-2:20','Eligibility report and scheme directory','“This is not a loan promise. It is an explainable profile, a source-cited scheme shortlist and the next best action.”'),
   ('2:20-3:05','Fraud/review console','“Signals raise a review flag. They do not automatically label a worker as fraudulent.”'),
   ('3:05-3:40','Architecture diagram','“Workers get simple guidance; lenders receive structured, auditable cases; consent and security run across every layer.”'),
   ('3:40-4:00','Vision and close','“Sahayata helps turn irregular digital income into understandable, responsibly assessed financial access.”')], [850,1900,6610])
  add_heading(doc,'Strong answers to likely judge questions',2)
  add_bullets(doc,['“How do you work without CIBIL?” — We do not replace bureau data. We help build a consent-based affordability profile for thin-file customers and route it through a lender’s regulated policy.','“Does AI approve loans?” — No. AI supports collection, explanation and risk prioritisation; licensed lenders own final approval, pricing and disbursal.','“How do you avoid fraud?” — We combine document, transaction, device and partner-verification signals with human review and an audit trail.','“How do workers understand schemes?” — The app gives multilingual, source-cited summaries with effective dates and directs users to official channels for final confirmation.'])
  # risks
  page_break(doc); add_heading(doc,'11. Risks, boundaries and compliance readiness',1); doc.add_picture(str(trust),width=Inches(6.85)); add_caption(doc,'Figure 6. The non-negotiable trust controls that make the system credible to workers, judges and lending partners.')
  add_table(doc,['Risk','Control / boundary'],[
   ('Privacy and consent','Use explicit purpose-based consent, data minimisation, encryption, access logs, retention schedule and revocation handling.'),
   ('Credit-model harm','Decision support only; independent validation, explainability, monitoring and human escalation before real lending use.'),
   ('Misleading claims','Never advertise guaranteed approval, fixed approval speed or savings without substantiated, current evidence.'),
   ('Scheme-rule changes','Versioned knowledge base with official source, effective date, owner and periodic review.'),
   ('Partner and API failure','Retries, timeout handling, safe pause states, reconciliation and incident runbooks.'),
   ('Fraud false positives','Confidence thresholds, appeal/correction path, investigator review and outcomes feedback loop.'),
   ('Regulatory scope','Engage regulated lenders and legal/compliance advisers before live AA, KYC, credit-bureau or collection integrations.')], [2350,7010])
  add_heading(doc,'Prototype status statement',2)
  add_para(doc,'The current product should be presented as a prototype and decision-support demonstration. Any live integration with Account Aggregators, Aadhaar/KYC rails, credit bureaus, lenders, mandates, payment systems or government portals must be launched only with authorised partners, contracts, security testing and applicable compliance approvals.',size=11,color=RED,bold=True)
  # sources
  page_break(doc); add_heading(doc,'12. Research sources and data notes',1)
  add_para(doc,'All statistics below are cited as reported by the named primary or institutional source. Dates matter: they should be rechecked before a live pitch or public release.',size=10,color=GRAY)
  sources=[
  ('ILO, India Employment Report 2024','Nearly 82% of workforce in informal sector and nearly 90% informally employed.','https://www.ilo.org/sites/default/files/2024-08/India%20Employment%20-%20web_8%20April.pdf'),
  ('Ministry of Labour & Employment / PIB, March 2025','Over 30.68 crore unorganised workers registered on e-Shram.','https://labour.gov.in/sites/default/files/pib2109828.pdf'),
  ('TransUnion CIBIL newsroom, 2022','35 million consumers opened a first credit product in 2021; 31 million in first nine months of 2022; 67% rural/semi-urban in 2021.','https://newsroom.transunioncibil.com/women-farmers-and-youth-lead-the-new-to-credit-consumer-segment-and-emerge-as-catalysts-of-sustainable-financial-inclusion-in-india/'),
  ('RBI / National Strategy for Financial Education 2020-2025','NCFE 2019 survey: 27.18% achieved the minimum target across all financial-literacy components; sample 75,000 adults.','https://systemhealth.rbi.org.in/Scripts/PublicationReportDetails.aspx_UrlPage%3D%26ID%3D1156%281%29.html'),
  ('RBI / PIB, August 2025','Financial Inclusion Index: 67.0 for year ending March 2025 versus 64.2 in March 2024.','https://www.pib.gov.in/PressNoteDetails.aspx?ModuleId=3&NoteId=154980&lang=1&reg=1'),
  ('NPCI UPI product statistics','Official month-by-month UPI transaction volume and value; use current page at presentation time.','https://www.npci.org.in/product/upi/product-statistics'),
  ('RBI financial-literacy and consumer-protection material','National financial-literacy initiatives and FAME messages.','https://www.rbi.org.in/commonperson/images/FAME202426022024.pdf'),
  ('PM SVANidhi official portal','Scheme details, eligibility and official process.','https://pmsvanidhi.mohua.gov.in/'),
  ('PM-SYM official portal','Unorganised-worker pension scheme information.','https://maandhan.in/'),
  ('Department of Financial Services, PMJDY','Financial-inclusion and PMJDY information.','https://pmjdy.gov.in/')]
  add_table(doc,['Source','Key use in this document','URL'],sources,[1900,3300,4160])
  add_heading(doc,'Data integrity note',2)
  add_para(doc,'No authoritative nationwide statistic was found that counts exactly “how many Indians have no CIBIL score.” This document intentionally does not invent one. The documented “new-to-credit” figures are a defensible, clearly labelled proxy for the thin-credit-file problem. Likewise, the financial-literacy survey measures capability—not direct awareness of every government scheme. This distinction makes the pitch more credible with judges.')
  alt_texts = [
    'Four India financial-inclusion context indicators: informal employment, e-Shram registrations, new-to-credit consumers and financial-literacy thresholds.',
    'Four connected Sahayata modules: worker experience, SAI assistant, trusted decisioning and lender operations.',
    'Six-step worker journey from SAI conversation through consent, eligibility, review, lender application and flexible repayment.',
    'Mermaid-style Sahayata production architecture showing worker, officer and regulated partner interactions with platform layers.',
    'AI-assisted technical flow from consent and evidence through safety gates to lender action and audit trail.',
    'Four trust controls: consent, security, explainability and accountability.'
  ]
  for shape, alt in zip(doc.inline_shapes, alt_texts):
    shape._inline.docPr.set('descr', alt)
    shape._inline.docPr.set('title', 'Sahayata diagram')
  doc.core_properties.title='Sahayata Judge Pitch and Technical Documentation'; doc.core_properties.author='Sahayata Team'; doc.save(DOCX)
  print(DOCX)

if __name__=='__main__': build()
